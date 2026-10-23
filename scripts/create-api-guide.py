from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

OUT = r'C:\Users\rafae\OneDrive\Documentos\ChatGPT\familyhub\docs\familyhub-guia-apis-credenciais.docx'
doc = Document()
section = doc.sections[0]
section.top_margin = Inches(.65)
section.bottom_margin = Inches(.65)
section.left_margin = Inches(.75)
section.right_margin = Inches(.75)
normal = doc.styles['Normal']
normal.font.name = 'Aptos'
normal.font.size = Pt(10)
normal.font.color.rgb = RGBColor(38, 43, 60)
for name, size, color in [('Title', 26, '20263A'), ('Heading 1', 17, '4338CA'), ('Heading 2', 12, '20263A')]:
    style = doc.styles[name]
    style.font.name = 'Aptos Display'
    style.font.size = Pt(size)
    style.font.bold = True
    style.font.color.rgb = RGBColor.from_string(color)

def shade(cell, fill):
    props = cell._tc.get_or_add_tcPr()
    element = OxmlElement('w:shd')
    element.set(qn('w:fill'), fill)
    props.append(element)

def text(value, style=None, bold=False, color=None):
    p = doc.add_paragraph(style=style)
    run = p.add_run(value)
    run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    return p

def bullets(items):
    for item in items:
        doc.add_paragraph(item, style='List Bullet')

def code(value):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(.18)
    run = p.add_run(value)
    run.font.name = 'Consolas'
    run.font.size = Pt(9)
    props = p._p.get_or_add_pPr()
    element = OxmlElement('w:shd')
    element.set(qn('w:fill'), 'F0F2FA')
    props.append(element)

def table(rows):
    table = doc.add_table(rows=0, cols=len(rows[0]))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = 'Table Grid'
    for index, row in enumerate(rows):
        cells = table.add_row().cells
        for cell, value in zip(cells, row):
            cell.text = value
            if index == 0:
                shade(cell, '4338CA')
            for run in cell.paragraphs[0].runs:
                run.font.size = Pt(9)
                run.font.bold = index == 0
                if index == 0:
                    run.font.color.rgb = RGBColor(255, 255, 255)

p = doc.add_paragraph(style='Title')
p.alignment = WD_ALIGN_PARAGRAPH.LEFT
p.add_run('FamilyHub Guia de APIs e Credenciais')
text('Configuração para uso pessoal da família', bold=True, color='4338CA')
text('Este documento explica quais serviços externos o FamilyHub utiliza, como criar as credenciais, onde colocá-las e quais informações não devem ser compartilhadas. Para o uso pessoal da sua família, a configuração pode permanecer local e sem contratação de serviços pagos.')

text('Resumo da decisão', style='Heading 1')
text('Você não precisa contratar APIs de terceiros para o funcionamento básico do FamilyHub. PostgreSQL, MinIO, Mailpit e Redis rodam localmente via Docker. A única integração externa prevista nesta etapa é o Google Agenda, que exige um projeto gratuito no Google Cloud e autorização explícita de cada conta Google.')
text('1 Serviços utilizados', style='Heading 1')
table([['Serviço', 'Finalidade', 'Credencial', 'Custo'], ['PostgreSQL', 'Banco principal', 'DATABASE_URL', 'Local e gratuito'], ['MinIO', 'Fotos e documentos privados', 'MINIO_ACCESS_KEY e MINIO_SECRET_KEY', 'Local e gratuito'], ['Mailpit', 'Testar e-mails', 'SMTP local', 'Local e gratuito'], ['Redis', 'Filas e jobs', 'REDIS_URL', 'Local e gratuito'], ['Google Calendar API', 'Enviar eventos autorizados', 'OAuth Client ID e Secret', 'Sem custo para uso pessoal normal']])

text('2 Serviços locais', style='Heading 1')
text('Na raiz do projeto, copie os arquivos de ambiente e inicialize o Docker Compose:')
code('Copy-Item apps/api/.env.example apps/api/.env\nCopy-Item apps/web/.env.example apps/web/.env\ndocker compose up -d\nnpm run db:deploy')
text('O Mailpit fica disponível em http://localhost:8025. Ele recebe os e-mails de confirmação, recuperação de senha e convites sem enviar mensagens para destinatários reais.')

text('3 Google Calendar API', style='Heading 1')
text('O Google Calendar usa OAuth 2.0. O Client ID e o Client Secret identificam o aplicativo FamilyHub; eles não representam uma família específica. Depois, cada administrador autoriza sua própria conta Google e a conexão fica vinculada àquela família.')
text('Criar o projeto', style='Heading 2')
bullets(['Acesse https://console.cloud.google.com/.', 'Crie um projeto chamado FamilyHub.', 'Abra APIs e serviços, depois Biblioteca.', 'Pesquise Google Calendar API e clique em Ativar.'])
text('Configurar consentimento', style='Heading 2')
bullets(['Abra Google Auth Platform e configure o nome do aplicativo e o e-mail de suporte.', 'Durante o desenvolvimento, adicione seu Gmail em Test users.', 'Mantenha o aplicativo em modo de teste para uso pessoal.'])
text('Criar o cliente OAuth', style='Heading 2')
bullets(['Abra Google Auth Platform e depois Clients.', 'Clique em Create Client.', 'Selecione Web application.', 'Em Authorized redirect URIs, adicione exatamente a URL abaixo.'])
code('http://localhost:3333/api/v1/integrations/google-calendar/callback')
text('O Google exibirá o Client ID e o Client Secret. Copie-os para apps/api/.env. Não use API Key para esse fluxo: a integração precisa de OAuth para obter a permissão da conta que possui o calendário.')
text('Gerar a chave de criptografia', style='Heading 2')
code("node -e \"console.log(require('crypto').randomBytes(32).toString('base64url'))\"")
text('Copie o resultado para INTEGRATION_ENCRYPTION_KEY. Essa chave protege os tokens OAuth no banco. Se ela for perdida, as contas Google precisarão ser reconectadas.')
text('Configuração final', style='Heading 2')
code('GOOGLE_CLIENT_ID=seu-client-id.apps.googleusercontent.com\nGOOGLE_CLIENT_SECRET=seu-client-secret\nGOOGLE_REDIRECT_URI=http://localhost:3333/api/v1/integrations/google-calendar/callback\nINTEGRATION_ENCRYPTION_KEY=chave-gerada\nAPP_URL=http://localhost:5173')

text('4 Uso por várias famílias', style='Heading 1')
text('As credenciais do Google Cloud são compartilhadas pelo aplicativo, mas a autorização é individual por família:')
code('FamilyHub\n  ├── Família A → Conta Google A → Calendário escolhido\n  ├── Família B → Conta Google B → Calendário escolhido\n  └── Família C → Conta Google A → Outro calendário escolhido')
bullets(['O administrador abre Família → Google Agenda.', 'Clica em Conectar Google Agenda.', 'Autoriza a conta Google.', 'Escolhe um calendário gravável.', 'Eventos criados no FamilyHub são enviados para esse calendário.'])
text('A primeira versão sincroniza somente FamilyHub → Google Agenda. Alterações feitas diretamente no Google ainda não retornam ao FamilyHub.')

text('5 Privacidade e segurança', style='Heading 1')
text('A integração envia título, descrição, local, horário e data dos eventos autorizados. Dados de saúde, perfis de crianças, documentos e financeiro não são sincronizados.')
bullets(['Nunca publique apps/api/.env no GitHub.', 'Nunca envie GOOGLE_CLIENT_SECRET ou INTEGRATION_ENCRYPTION_KEY pelo chat.', 'Use HTTPS e uma redirect URI HTTPS quando o sistema sair do ambiente local.', 'Se uma chave for exposta, revogue o cliente OAuth no Google Cloud e gere outro.'])
table([['Segredo', 'Proteção'], ['GOOGLE_CLIENT_SECRET', 'Nunca vai para o front-end ou GitHub.'], ['INTEGRATION_ENCRYPTION_KEY', 'Protege tokens Google armazenados no banco.'], ['JWT_SECRET', 'Protege sessões do FamilyHub.'], ['MINIO_SECRET_KEY', 'Protege o armazenamento local.'], ['DATABASE_URL', 'Contém acesso ao banco PostgreSQL.']])

text('6 Como testar', style='Heading 1')
bullets(['Execute npm run dev.', 'Entre como administrador da família.', 'Abra Família → Google Agenda.', 'Conecte a conta Google e escolha o calendário.', 'Crie um evento no Calendário do FamilyHub.', 'Confirme o evento no Google Agenda.', 'Use Sincronizar agora se necessário.', 'Abra http://localhost:8025 para testar os e-mails.'])
text('7 Quando não configurar APIs externas', style='Heading 1')
text('Se você não quiser usar Google Agenda, não precisa configurar credenciais Google. Login, famílias, tarefas, compras, calendário interno, perfil, e-mails locais e armazenamento local continuam funcionando com Docker e os arquivos .env padrão.')
text('Referências', style='Heading 1')
text('Google Calendar API: https://developers.google.com/workspace/calendar/api\nOAuth para aplicações web: https://developers.google.com/identity/protocols/oauth2/web-server\nGuia específico do projeto: docs/google-calendar.md')
doc.core_properties.title = 'FamilyHub Guia de APIs e Credenciais'
doc.core_properties.subject = 'Configuração de APIs para uso pessoal'
doc.core_properties.author = 'FamilyHub'
doc.save(OUT)
print(OUT)
